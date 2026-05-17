from flask import Blueprint, request, jsonify
from datetime import datetime
import openpyxl
import io

from models import get_db, validate_israeli_phone

guests_bp = Blueprint("guests", __name__)


def row_to_dict(row):
    return dict(row)


@guests_bp.route("", methods=["GET"])
def list_guests():
    side = request.args.get("side")
    grp = request.args.get("grp")
    status = request.args.get("status")
    search = request.args.get("search", "").strip()

    query = "SELECT * FROM guests WHERE 1=1"
    params = []

    if side:
        query += " AND side = ?"
        params.append(side)
    if grp:
        query += " AND grp = ?"
        params.append(grp)
    if status:
        query += " AND rsvp_status = ?"
        params.append(status)
    if search:
        query += " AND name LIKE ?"
        params.append(f"%{search}%")

    query += " ORDER BY created_at DESC"

    conn = get_db()
    rows = conn.execute(query, params).fetchall()
    conn.close()
    return jsonify([row_to_dict(r) for r in rows])


@guests_bp.route("/<int:guest_id>", methods=["GET"])
def get_guest(guest_id):
    conn = get_db()
    row = conn.execute("SELECT * FROM guests WHERE id = ?", (guest_id,)).fetchone()
    conn.close()
    if not row:
        return jsonify({"error": "לא נמצא"}), 404
    return jsonify(row_to_dict(row))


@guests_bp.route("", methods=["POST"])
def add_guest():
    data = request.get_json()
    required = ["name", "side", "grp"]
    for field in required:
        if not data.get(field):
            return jsonify({"error": f"שדה חסר: {field}"}), 400

    phone = data.get("phone", "")
    valid = 1 if validate_israeli_phone(phone) else 0

    conn = get_db()
    cur = conn.execute("""
        INSERT INTO guests (name, phone, email, side, grp, party_size, notes,
                            rsvp_status, phone_valid, created_at)
        VALUES (?,?,?,?,?,?,?,?,?,?)
    """, (
        data["name"].strip(),
        phone,
        data.get("email", ""),
        data["side"],
        data["grp"],
        int(data.get("party_size", 1)),
        data.get("notes", ""),
        data.get("rsvp_status", "pending"),
        valid,
        datetime.now().isoformat(),
    ))
    conn.commit()
    row = conn.execute("SELECT * FROM guests WHERE id = ?", (cur.lastrowid,)).fetchone()
    conn.close()
    return jsonify(row_to_dict(row)), 201


@guests_bp.route("/<int:guest_id>", methods=["PUT"])
def update_guest(guest_id):
    data = request.get_json()
    conn = get_db()
    existing = conn.execute("SELECT * FROM guests WHERE id = ?", (guest_id,)).fetchone()
    if not existing:
        conn.close()
        return jsonify({"error": "לא נמצא"}), 404

    phone = data.get("phone", existing["phone"]) or ""
    valid = 1 if validate_israeli_phone(phone) else 0

    # If status changed to attending/not_attending and wasn't before, set responded_at
    old_status = existing["rsvp_status"]
    new_status = data.get("rsvp_status", old_status)
    responded_at = existing["responded_at"]
    if old_status == "pending" and new_status != "pending" and not responded_at:
        responded_at = datetime.now().isoformat()

    conn.execute("""
        UPDATE guests SET
            name        = ?,
            phone       = ?,
            email       = ?,
            side        = ?,
            grp         = ?,
            party_size  = ?,
            notes       = ?,
            rsvp_status = ?,
            phone_valid = ?,
            responded_at = ?
        WHERE id = ?
    """, (
        data.get("name", existing["name"]),
        phone,
        data.get("email", existing["email"]),
        data.get("side", existing["side"]),
        data.get("grp", existing["grp"]),
        int(data.get("party_size", existing["party_size"])),
        data.get("notes", existing["notes"]),
        new_status,
        valid,
        responded_at,
        guest_id,
    ))
    conn.commit()
    row = conn.execute("SELECT * FROM guests WHERE id = ?", (guest_id,)).fetchone()
    conn.close()
    return jsonify(row_to_dict(row))


@guests_bp.route("/<int:guest_id>", methods=["DELETE"])
def delete_guest(guest_id):
    conn = get_db()
    existing = conn.execute("SELECT id FROM guests WHERE id = ?", (guest_id,)).fetchone()
    if not existing:
        conn.close()
        return jsonify({"error": "לא נמצא"}), 404
    conn.execute("DELETE FROM guests WHERE id = ?", (guest_id,))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@guests_bp.route("/bulk-delete", methods=["POST"])
def bulk_delete():
    data = request.get_json()
    ids = data.get("ids", [])
    if not ids:
        return jsonify({"error": "אין מזהים"}), 400
    placeholders = ",".join("?" * len(ids))
    conn = get_db()
    conn.execute(f"DELETE FROM guests WHERE id IN ({placeholders})", ids)
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "deleted": len(ids)})


@guests_bp.route("/import/preview", methods=["POST"])
def import_preview():
    """Parse uploaded .xlsx and return rows for preview (not saved yet)."""
    if "file" not in request.files:
        return jsonify({"error": "קובץ חסר"}), 400

    file = request.files["file"]
    if not file.filename.endswith((".xlsx", ".xls")):
        return jsonify({"error": "יש להעלות קובץ Excel בלבד (.xlsx)"}), 400

    content = file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        return jsonify({"error": f"שגיאה בקריאת הקובץ: {str(e)}"}), 400

    if not rows:
        return jsonify({"error": "הקובץ ריק"}), 400

    # Detect header row
    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    col_map = {
        "name": next((i for i, h in enumerate(headers) if "שם" in h or "name" in h), None),
        "phone": next((i for i, h in enumerate(headers) if "טלפון" in h or "phone" in h or "נייד" in h), None),
        "email": next((i for i, h in enumerate(headers) if "מייל" in h or "email" in h), None),
        "side": next((i for i, h in enumerate(headers) if "צד" in h or "side" in h), None),
        "grp": next((i for i, h in enumerate(headers) if "קבוצה" in h or "group" in h or "grp" in h), None),
        "party_size": next((i for i, h in enumerate(headers) if "כמות" in h or "party" in h or "size" in h or "אורחים" in h), None),
        "notes": next((i for i, h in enumerate(headers) if "הערות" in h or "notes" in h), None),
    }

    SIDE_MAP = {"כלה": "bride", "חתן": "groom", "bride": "bride", "groom": "groom"}
    GROUP_MAP = {"משפחה": "family", "חברים": "friends", "עבודה": "work", "אחר": "other",
                 "family": "family", "friends": "friends", "work": "work", "other": "other"}

    preview = []
    errors = []
    for i, row in enumerate(rows[1:], start=2):
        def cell(key):
            idx = col_map.get(key)
            if idx is None or idx >= len(row):
                return ""
            val = row[idx]
            return str(val).strip() if val is not None else ""

        name = cell("name")
        if not name:
            errors.append(f"שורה {i}: שם חסר, דולגה")
            continue

        phone = cell("phone")
        raw_side = cell("side").lower()
        raw_grp = cell("grp").lower()

        entry = {
            "name": name,
            "phone": phone,
            "email": cell("email"),
            "side": SIDE_MAP.get(raw_side, "bride"),
            "grp": GROUP_MAP.get(raw_grp, "other"),
            "party_size": int(cell("party_size") or 1),
            "notes": cell("notes"),
            "phone_valid": validate_israeli_phone(phone),
            "rsvp_status": "pending",
        }
        preview.append(entry)

    return jsonify({"preview": preview, "errors": errors, "col_map": col_map})


@guests_bp.route("/import/confirm", methods=["POST"])
def import_confirm():
    """Save previewed guests to DB."""
    data = request.get_json()
    guests_data = data.get("guests", [])
    if not guests_data:
        return jsonify({"error": "אין נתונים לשמירה"}), 400

    conn = get_db()
    inserted = 0
    now = datetime.now().isoformat()
    for g in guests_data:
        phone = g.get("phone", "")
        valid = 1 if validate_israeli_phone(phone) else 0
        conn.execute("""
            INSERT INTO guests (name, phone, email, side, grp, party_size, notes,
                                rsvp_status, phone_valid, created_at)
            VALUES (?,?,?,?,?,?,?,?,?,?)
        """, (
            g.get("name", ""),
            phone,
            g.get("email", ""),
            g.get("side", "bride"),
            g.get("grp", "other"),
            int(g.get("party_size", 1)),
            g.get("notes", ""),
            "pending",
            valid,
            now,
        ))
        inserted += 1
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "inserted": inserted})
